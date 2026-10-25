"""Override audit: how often the final pipeline care_type differs from the
raw ML argmax, which stage changed it, and whether each override helped or
hurt accuracy.

Runs three sets through the full pipeline (offline: the LLM is disabled, so
only the deterministic stages run — the LLM second opinion never changes
care_type, it can only ask a confirmation question):

  1. ml/data/unseen_tests.csv            (30 rows, raw = production model argmax)
  2. ml/data/redflag_holdout.csv         (53 rows, raw = production model argmax)
  3. official triage_dataset.jsonl       (raw = exact-text grouped CV
     out-of-fold prediction, the honest "raw model on its own" estimate
     recomputed exactly as in compare_models.py: StratifiedGroupKFold on
     exact normalized text, seed 42, TF-IDF char (2-5) + LR C=2.0 balanced)

No thresholds are tuned here. Outputs:
  - console report
  - ml/results/override_audit.json
  - ml/results/override_audit_rows.csv (per-row trail)
  - docs/verification_workbook.xlsx (Summary sheet updated)
"""
import csv
import json
import sys
from collections import Counter
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.config import settings

settings.DEEPSEEK_API_KEY = ""  # offline: LLM second opinion never changes care_type anyway

from app.schemas import TriageRequest, Location
from app.core.pipeline import run_triage_pipeline
from app.core.triage_model import raw_model_argmax
from app.core.language import normalise_text
from app.core.intake import extract_slots, build_classifier_text
from app.core.guardrails import input_guard

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
RESULTS_DIR = BASE_DIR / "results"
RESULTS_DIR.mkdir(parents=True, exist_ok=True)
DOCS_DIR = BASE_DIR.parent / "docs"

UNSEEN = DATA_DIR / "unseen_tests.csv"
HOLDOUT = DATA_DIR / "redflag_holdout.csv"
OFFICIAL = DATA_DIR / "official" / "triage_dataset.jsonl"

LOC = Location(lat=28.54, lng=77.27)

_LANG_MAP = {"English": "en", "Hindi": "hi", "Hinglish": "hinglish",
             "en": "en", "hi": "hi", "hinglish": "hinglish"}

STAGE_LABELS = {
    "5_hybrid_rules_ml": "Stage 5 rule layer (inside predict_triage)",
    "5_unfamiliar_map": "Stage 5 unfamiliar-care-type mapping",
    "6_risk_modifiers": "Stage 6 risk modifiers",
    "6.5_ambulatory": "Stage 6.5 ambulatory adjustment",
    "6.6_injury_sanity": "Stage 6.6 injury sanity",
    "7_confidence_fallback": "Stage 7 confidence fallback",
    "8_emergency_rule2": "Stage 8 emergency rule (Rule 2)",
    "8_emergency_remap": "Stage 8 emergency remap (not-emergency)",
    "8_pregnancy_floor": "Stage 8 pregnancy floor",
    "8_context_night_rain": "Stage 8 context night/rain",
}


def load_official_clf_entries():
    entries = []
    with open(OFFICIAL, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            item = json.loads(line)
            req = item.get("request", {})
            resp = item.get("response", {})
            raw = (req.get("user_input") or "").strip()
            norm = normalise_text(raw)
            if not norm:
                continue
            entries.append({
                "id": item.get("id"),
                "user_input": raw,
                "norm_text": norm,
                "language": item.get("detected_language") or req.get("language") or "English",
                "care_type": resp.get("care_type"),
            })
    return [e for e in entries if e.get("care_type")]


def exact_text_cv_oof(entries):
    """Replicates compare_models.py's exact-text grouped CV for Model A and
    returns the out-of-fold prediction per entry (the honest raw-model
    estimate on training data)."""
    texts = [e["norm_text"] for e in entries]
    labels = [e["care_type"] for e in entries]
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.linear_model import LogisticRegression
    from sklearn.model_selection import StratifiedGroupKFold

    text_to_grp = {t: i for i, t in enumerate(sorted(set(texts)))}
    exact_groups = np.array([text_to_grp[t] for t in texts])
    sgkf = StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=42)
    oof = np.empty(len(texts), dtype=object)
    for tr, val in sgkf.split(texts, labels, groups=exact_groups):
        vec = TfidfVectorizer(ngram_range=(2, 5), analyzer="char", min_df=1, sublinear_tf=True)
        x_tr = vec.fit_transform([texts[i] for i in tr])
        x_v = vec.transform([texts[i] for i in val])
        clf = LogisticRegression(C=2.0, class_weight="balanced", max_iter=500, random_state=42)
        clf.fit(x_tr, [labels[i] for i in tr])
        oof[val] = clf.predict(x_v)
    return [str(p) for p in oof]


def replicated_raw_argmax(user_input):
    """Raw production-model argmax computed exactly the way the pipeline
    builds its classifier text (used for rows that emergency-bypass stage 5,
    where the pipeline itself never ran the model)."""
    cleaned = input_guard(user_input)["cleaned_text"]
    norm = normalise_text(cleaned)
    slots = extract_slots(norm, [])
    return raw_model_argmax(build_classifier_text(slots))


def run_row(row):
    """Runs the pipeline for one row; returns (final, overrides, raw_prod)."""
    lang = _LANG_MAP.get(row.get("language", "en"), "en")
    resp = run_triage_pipeline(
        TriageRequest(
            user_input=row["user_input"],
            language=lang,
            location=LOC,
            conversation_history=[],
        ),
        debug=True,
    )
    debug = resp.debug or {}
    overrides = debug.get("overrides") or []
    raw_prod = None
    for o in overrides:
        if o.get("stage") == "5_model_argmax":
            raw_prod = o.get("care_type")
    if raw_prod is None and resp.care_type is not None:
        # Emergency bypass: stage 5 never ran; replicate the classifier text.
        raw_prod = replicated_raw_argmax(row["user_input"])
    return resp.care_type, overrides, raw_prod


def attributing_stage(overrides, final, raw):
    """The stage responsible for the final deviation from the raw argmax."""
    for o in reversed(overrides):
        if o.get("to") == final:
            return o.get("stage")
    if final == raw:
        return None
    # The stage-5 hybrid already equals the final: the rule layer inside
    # predict_triage produced it.
    return "5_hybrid_rules_ml"


def evaluate(name, rows, cv_oof=None):
    stats = {
        "name": name,
        "n_rows": len(rows),
        "n_decided": 0,
        "n_undecided": 0,
        "n_labeled_decided": 0,
        "n_override": 0,
        "n_override_labeled": 0,
        "override_fixed": 0,
        "override_broke": 0,
        "override_both_wrong": 0,
        "pipeline_correct": 0,
        "raw_correct": 0,
        "stage_transitions": Counter(),
        "stage_attribution": Counter(),
        "rows": [],
    }
    for i, row in enumerate(rows):
        final, overrides, raw_prod = run_row(row)
        if cv_oof is not None:
            raw = cv_oof[i]
        else:
            raw = raw_prod
        if final is None:
            stats["n_undecided"] += 1
            stats["rows"].append({
                "id": row.get("id"), "user_input": row["user_input"],
                "raw_ml": raw, "final": None, "true_label": row.get("care_type") or "",
                "override": "", "stage": "", "outcome": "",
            })
            continue
        stats["n_decided"] += 1
        for o in overrides:
            if o.get("from") and o.get("to") and o.get("from") != o.get("to"):
                stats["stage_transitions"][o.get("stage")] += 1
        override = bool(raw is not None and final != raw)
        stage = attributing_stage(overrides, final, raw) if override else ""
        outcome = ""
        if override:
            stats["n_override"] += 1
            stats["stage_attribution"][stage] += 1
        true = row.get("care_type")
        if true:
            stats["n_labeled_decided"] += 1
            if final == true:
                stats["pipeline_correct"] += 1
            if raw is not None and raw == true:
                stats["raw_correct"] += 1
            if override:
                stats["n_override_labeled"] += 1
                if final == true:
                    stats["override_fixed"] += 1
                    outcome = "fixed"
                elif raw == true:
                    stats["override_broke"] += 1
                    outcome = "broke"
                else:
                    stats["override_both_wrong"] += 1
                    outcome = "both_wrong"
        stats["rows"].append({
            "id": row.get("id"), "user_input": row["user_input"],
            "raw_ml": raw or "", "final": final, "true_label": true or "",
            "override": "1" if override else "", "stage": stage, "outcome": outcome,
        })
    stats["override_rate"] = (stats["n_override"] / stats["n_decided"]) if stats["n_decided"] else None
    stats["pipeline_acc"] = (stats["pipeline_correct"] / stats["n_labeled_decided"]) if stats["n_labeled_decided"] else None
    stats["raw_acc"] = (stats["raw_correct"] / stats["n_labeled_decided"]) if stats["n_labeled_decided"] else None
    return stats


def update_workbook(report, stage_labels):
    """Writes the Summary sheet of docs/verification_workbook.xlsx."""
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        print("openpyxl not installed; skipping verification workbook update")
        return
    wb_path = DOCS_DIR / "verification_workbook.xlsx"
    if wb_path.exists():
        wb = load_workbook(wb_path)
        if "Summary" in wb.sheetnames:
            ws = wb["Summary"]
        else:
            ws = wb.create_sheet("Summary", 0)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
    ws["A1"] = "CareCompass Verification Summary"
    ws["A2"] = "Override audit: final care_type vs raw ML argmax (no thresholds tuned)"
    row = 4
    for col, head in enumerate(["Dataset", "Rows", "Decided", "Overridden", "Override rate",
                                "Pipeline acc", "Raw ML acc"], start=1):
        ws.cell(row=row, column=col, value=head)
    row += 1
    for d in report["datasets"]:
        ws.cell(row=row, column=1, value=d["name"])
        ws.cell(row=row, column=2, value=d["n_rows"])
        ws.cell(row=row, column=3, value=d["n_decided"])
        ws.cell(row=row, column=4, value=d["n_override"])
        ws.cell(row=row, column=5, value=round(d["override_rate"], 4))
        ws.cell(row=row, column=6, value=round(d["pipeline_acc"], 4))
        ws.cell(row=row, column=7, value=round(d["raw_acc"], 4))
        row += 1
    ws.cell(row=row, column=1, value="ALL")
    ws.cell(row=row, column=3, value=report["all"]["n_decided"])
    ws.cell(row=row, column=4, value=report["all"]["n_override"])
    ws.cell(row=row, column=5, value=round(report["all"]["override_rate"], 4))
    row += 2
    ws.cell(row=row, column=1, value="Stage")
    ws.cell(row=row, column=2, value="Meaning")
    ws.cell(row=row, column=3, value="Transitions")
    ws.cell(row=row, column=4, value="Attributed overrides")
    row += 1
    for k, label in stage_labels.items():
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=label)
        ws.cell(row=row, column=3, value=report["stages"][k]["transitions"])
        ws.cell(row=row, column=4, value=report["stages"][k]["attributed_overrides"])
        row += 1
    wb.save(wb_path)
    print(f"Updated {wb_path} (Summary sheet)")


def main():
    unseen = []
    with open(UNSEEN, "r", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            unseen.append({"id": r["id"], "user_input": r["user_input"],
                           "language": r.get("language", "en"), "care_type": r["care_type"]})
    holdout = []
    with open(HOLDOUT, "r", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            holdout.append({"id": r["id"], "user_input": r["user_input"],
                            "language": r.get("language", "en"), "care_type": r["care_type"]})
    official = load_official_clf_entries()
    print(f"Official rows (labeled, exact-text CV set): {len(official)}")
    cv_oof = exact_text_cv_oof(official)

    datasets = [
        ("unseen_tests.csv", unseen, None),
        ("redflag_holdout.csv", holdout, None),
        ("official_triage (exact-text CV OOF raw)", official, cv_oof),
    ]

    print("\nRunning pipeline (offline, deterministic stages only)...\n")
    report = {"datasets": [], "stages": {}, "all": {}}
    all_decided = 0
    all_override = 0
    transitions_total = Counter()
    attribution_total = Counter()
    rows_all = []
    for name, rows, oof in datasets:
        print(f"  {name}: {len(rows)} rows")
        stats = evaluate(name, rows, cv_oof=oof)
        report["datasets"].append(stats)
        rows_all.extend((name, r) for r in stats["rows"])
        all_decided += stats["n_decided"]
        all_override += stats["n_override"]
        transitions_total.update(stats["stage_transitions"])
        attribution_total.update(stats["stage_attribution"])
        print(f"    decided={stats['n_decided']} undecided={stats['n_undecided']} "
              f"override_rate={stats['override_rate']:.3f} "
              f"pipeline_acc={stats['pipeline_acc']:.3f} raw_acc={stats['raw_acc']:.3f}")

    report["all"] = {
        "n_decided": all_decided,
        "n_override": all_override,
        "override_rate": (all_override / all_decided) if all_decided else None,
    }
    report["stages"] = {k: {"transitions": transitions_total.get(k, 0),
                            "attributed_overrides": attribution_total.get(k, 0)}
                        for k in STAGE_LABELS}

    print("\n" + "=" * 90)
    print("OVERRIDE AUDIT REPORT")
    print("=" * 90)
    print(f"\nOverall override rate: {all_override}/{all_decided} decided rows "
          f"({report['all']['override_rate']:.3f})")
    print("\nPer-dataset:")
    for d in report["datasets"]:
        print(f"  {d['name']:40s} rows={d['n_rows']:>4} decided={d['n_decided']:>4} "
              f"override={d['n_override']:>4} ({d['override_rate']:.3f}) "
              f"pipeline_acc={d['pipeline_acc']:.3f} raw_acc={d['raw_acc']:.3f}")
    print("\nPer-stage (transitions = rows where the stage changed care_type;")
    print("attributed = overrides whose final value came from that stage):")
    for k, label in STAGE_LABELS.items():
        print(f"  {k:26s} {label:44s} transitions={transitions_total.get(k, 0):>3} "
              f"attributed={attribution_total.get(k, 0):>3}")
    print("\nOverride correctness vs true label (labeled overrides):")
    for d in report["datasets"]:
        print(f"  {d['name']:40s} fixed={d['override_fixed']:>3} broke={d['override_broke']:>3} "
              f"both_wrong={d['override_both_wrong']:>3} (of {d['n_override_labeled']} labeled)")

    with open(RESULTS_DIR / "override_audit.json", "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, default=str)
    print(f"\nSaved ml/results/override_audit.json")

    with open(RESULTS_DIR / "override_audit_rows.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["dataset", "id", "user_input", "raw_ml", "final", "true_label",
                    "override", "attributed_stage", "override_outcome"])
        for ds_name, r in rows_all:
            w.writerow([ds_name, r["id"], r["user_input"], r["raw_ml"], r["final"] or "",
                        r["true_label"], r["override"], r["stage"], r["outcome"]])
    print("Saved ml/results/override_audit_rows.csv")

    update_workbook(report, STAGE_LABELS)


if __name__ == "__main__":
    main()
